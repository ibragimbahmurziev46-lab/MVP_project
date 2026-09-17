"""Экспорт материала в DOCX и PDF (разделы 9, 14 ТЗ).

DOCX — через python-docx (все платформы).
PDF — через WeasyPrint (Linux/server) или ReportLab (fallback).
"""

from __future__ import annotations

import io
import logging
import re
from html.parser import HTMLParser

from docx.shared import Pt

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# DOCX: HTML → python-docx (allowed tags only: h1–h3, p, strong, em,
#       ul/ol/li, table, br)
# ---------------------------------------------------------------------------

class _DocxHtmlParser(HTMLParser):
    """Разбирает допустимые HTML-теги и накапливает документ python-docx."""

    def __init__(self, doc):
        super().__init__(convert_charrefs=True)
        self.doc = doc
        self._list_type: str | None = None
        self._paragraph = None
        self._heading_level = 0
        self._bold = False
        self._italic = False
        self._segments: list[tuple[str, bool, bool]] = []

        self._in_cell = False
        self._cell_lines: list[str] = []
        self._row_cells: list[str] = []
        self._rows: list[list[str]] = []

    # helpers --------------------------------------------------------------

    def _flush_paragraph(self):
        if self._paragraph is None:
            self._segments = []
            return
        segments = self._segments
        self._segments = []
        text = "".join(s[0] for s in segments).strip()
        if not text:
            # убрать пустой абзац
            self._paragraph._element.getparent().remove(self._paragraph._element)
            self._paragraph = None
            return
        if len(segments) == 1 and segments[0][1] == segments[0][2] == False:
            self._paragraph.text = text
        else:
            for run_text, bold, italic in segments:
                run_text = run_text.strip()
                if not run_text:
                    continue
                run = self._paragraph.add_run(run_text)
                run.bold = bold
                run.italic = italic
        self._paragraph = None

    def _collect(self, data):
        if not data:
            return
        if self._in_cell:
            self._cell_lines.append(data)
            return
        if self._segments and self._segments[-1][1:] == (self._bold, self._italic):
            prev = self._segments[-1]
            self._segments[-1] = (prev[0] + data, prev[1], prev[2])
        else:
            self._segments.append((data, self._bold, self._italic))

    def _begin_paragraph(self):
        if self._heading_level:
            self._paragraph = self.doc.add_heading(level=self._heading_level)
        elif self._list_type:
            style = "List Bullet" if self._list_type == "ul" else "List Number"
            self._paragraph = self.doc.add_paragraph(style=style)
        else:
            self._paragraph = self.doc.add_paragraph()
        self._segments = []

    def _end_heading(self):
        self._flush_paragraph()
        self._heading_level = 0

    # Parser callbacks ------------------------------------------------------

    def handle_starttag(self, tag, attrs):
        if tag in ("h1", "h2", "h3"):
            self._flush_paragraph()
            self._heading_level = {"h1": 1, "h2": 2, "h3": 3}[tag]
            self._begin_paragraph()
        elif tag == "p":
            self._flush_paragraph()
            self._list_type = None
            self._begin_paragraph()
        elif tag == "br":
            self._segments.append(("\n", self._bold, self._italic))
        elif tag == "strong":
            self._bold = True
        elif tag == "em":
            self._italic = True
        elif tag in ("ul", "ol"):
            self._flush_paragraph()
            self._list_type = tag
            self._begin_paragraph()
        elif tag == "li":
            self._flush_paragraph()
            self._begin_paragraph()
        elif tag == "table":
            self._flush_paragraph()
            self._rows = []
        elif tag == "tr":
            self._row_cells = []
        elif tag in ("td", "th"):
            self._in_cell = True
            self._cell_lines = []

    def handle_endtag(self, tag):
        if tag in ("h1", "h2", "h3"):
            self._end_heading()
        elif tag == "p":
            self._flush_paragraph()
        elif tag == "strong":
            self._bold = False
        elif tag == "em":
            self._italic = False
        elif tag in ("ul", "ol"):
            self._flush_paragraph()
            self._list_type = None
        elif tag == "li":
            self._flush_paragraph()
        elif tag in ("td", "th"):
            self._in_cell = False
            self._row_cells.append(" ".join(self._cell_lines).strip())
        elif tag == "tr":
            if self._row_cells:
                self._rows.append(self._row_cells)
        elif tag == "table":
            self._build_table()

    def handle_data(self, data):
        self._collect(data)

    def _build_table(self):
        if not self._rows:
            return
        ncols = max(len(row) for row in self._rows)
        table = self.doc.add_table(rows=0, cols=ncols)
        table.style = "Table Grid"
        for row in self._rows:
            cells = table.add_row().cells
            for i, value in enumerate(row):
                if i < len(cells):
                    cells[i].text = value


def generate_docx(html: str, title: str = "Документ") -> io.BytesIO:
    """Принимает html и title, возвращает BytesIO с DOCX."""
    from docx import Document

    doc = Document()

    # Шрифт по умолчанию
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(14)

    parser = _DocxHtmlParser(doc)
    try:
        parser.feed(html)
        parser.close()
        parser._flush_paragraph()
    except Exception as exc:
        logger.exception("Ошибка конвертации HTML → DOCX")
        doc.add_paragraph(f"Ошибка экспорта: {exc}")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# PDF: сначала пытаемся WeasyPrint, если не установлен — ReportLab fallback
# ---------------------------------------------------------------------------

_PDF_FONT = "Times-Roman"
_PDF_BOLD = "Times-Bold"
_PDF_ITALIC = "Times-Italic"
_PDF_FONTSIZE = 14


class _PdfHtmlParser(HTMLParser):
    def __init__(self, pdf_doc):
        super().__init__(convert_charrefs=True)
        self.pdf_doc = pdf_doc
        self._bold = False
        self._italic = False
        self._heading = 0
        self._line_parts: list[str] = []
        self._paragraph_started = False

    def _start_line(self):
        if not self._paragraph_started:
            self._paragraph_started = True

    def _flush_line(self, style_tag=None):
        text = "".join(self._line_parts).strip()
        self._line_parts = []
        self._paragraph_started = False
        if not text:
            return

        if style_tag and style_tag in ("ul", "ol"):
            self.pdf_doc.add_paragraph("• " + text if style_tag == "ul" else text)
            return

        size = _PDF_FONTSIZE
        if self._heading:
            size = {1: 20, 2: 17, 3: 14}.get(self._heading, 12)
        self.pdf_doc.add_paragraph(text)

    def handle_starttag(self, tag, attrs):
        if tag in ("h1", "h2", "h3"):
            self._flush_line()
            self._heading = {"h1": 1, "h2": 2, "h3": 3}[tag]
        elif tag in ("p", "br"):
            self._flush_line()
            self._heading = 0
        elif tag == "strong":
            self._bold = True
        elif tag == "em":
            self._italic = True
        elif tag == "li":
            self._flush_line("ul" if self._heading else None)

    def handle_endtag(self, tag):
        if tag in ("h1", "h2", "h3"):
            self._flush_line()
            self._heading = 0
        elif tag == "p":
            self._flush_line()
        elif tag == "strong":
            self._bold = False
        elif tag == "em":
            self._italic = False
        elif tag == "li":
            self._flush_line("ul")

    def handle_data(self, data):
        self._line_parts.append(data)


def _generate_pdf_reportlab(html: str, title: str = "Документ") -> io.BytesIO:
    """PDF на ReportLab. Регистрирует системный TTF-шрифт для кириллицы."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph as RLParagraph, Spacer
    )
    from reportlab.lib.enums import TA_LEFT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    # Найти шрифт с кириллицей (Windows)
    import os
    font_candidates = [
        (r"C:\Windows\Fonts\times.ttf", "TNR"),
        (r"C:\Windows\Fonts\arial.ttf", "Arial"),
        (r"C:\Windows\Fonts\calibri.ttf", "Calibri"),
    ]
    font_name = _PDF_FONT
    for path, name in font_candidates:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont(name, path))
                font_name = name
                break
            except Exception:
                continue

    normal = ParagraphStyle(
        "Normal",
        fontName=font_name,
        fontSize=_PDF_FONTSIZE,
        leading=20,
        alignment=TA_LEFT,
    )
    heading = ParagraphStyle(
        "Heading",
        fontName=font_name,
        fontSize=16,
        leading=22,
        spaceBefore=12,
        alignment=TA_LEFT,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )

    clean = re.sub(r"<br\s*/?>", "\n", html)
    clean = re.sub(r"</?(table|thead|tbody|tr|th|td)[^>]*>", "\t", clean)
    clean = re.sub(r"<[^>]+>", "", clean).strip()
    clean = re.sub(r"\n+", "\n", clean)
    clean = re.sub(r"[ \t]+", " ", clean)

    story = []
    for para in clean.split("\n"):
        para = para.strip()
        if not para:
            continue
        style = heading if para.endswith(":") and len(para) < 80 else normal
        story.append(RLParagraph(para.replace("\t", " &nbsp; "), style))
        story.append(Spacer(1, 8))

    if not story:
        story.append(RLParagraph("Документ не сформирован.", normal))

    doc.build(story)
    buf.seek(0)
    return buf


def generate_pdf(html: str, title: str = "Документ") -> io.BytesIO:
    """Генерирует PDF из HTML."""
    try:
        import weasyprint  # noqa: F401
        doc = weasyprint.HTML(string=html).write_pdf()
        return io.BytesIO(doc)
    except ImportError:
        pass
    except Exception as exc:
        logger.warning("WeasyPrint ошибка (%s), переключаюсь на ReportLab", exc)

    try:
        return _generate_pdf_reportlab(html, title)
    except Exception as exc:
        logger.exception("ReportLab тоже не сработал")
        buf = io.BytesIO()
        buf.write(f"Ошибка экспорта PDF: {exc}".encode("utf-8"))
        buf.seek(0)
        return buf